#!/usr/bin/env python3
"""Read-only source extractor for the SDS188 freeze-gate audit."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_pdf(path: Path, out_dir: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"\n===== PAGE {idx} =====\n{text}")
    out_path = out_dir / f"{path.stem}.txt"
    out_path.write_text("".join(pages), encoding="utf-8")
    return {
        "path": str(path),
        "kind": "pdf",
        "pages": len(reader.pages),
        "encrypted": reader.is_encrypted,
        "text_output": str(out_path),
    }


def extract_docx(path: Path, out_dir: Path) -> dict:
    doc = Document(str(path))
    blocks = []
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        if paragraph.text.strip():
            blocks.append(f"P{idx}: {paragraph.text}")
    for t_idx, table in enumerate(doc.tables, start=1):
        blocks.append(f"\n===== TABLE {t_idx} =====")
        for r_idx, row in enumerate(table.rows, start=1):
            blocks.append(f"R{r_idx}: " + " | ".join(cell.text for cell in row.cells))
    out_path = out_dir / f"{path.stem}.txt"
    out_path.write_text("\n".join(blocks), encoding="utf-8")
    return {
        "path": str(path),
        "kind": "docx",
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "text_output": str(out_path),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("files", nargs="+", type=Path)
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for path in args.files:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            results.append(extract_pdf(path, args.out_dir))
        elif suffix == ".docx":
            results.append(extract_docx(path, args.out_dir))
        else:
            raise ValueError(f"Unsupported file: {path}")
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
