#!/usr/bin/env python3
"""Read-only OOXML/openpyxl inventory for workbook auditing."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


def workbook_report(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    wb_values = load_workbook(path, data_only=True, read_only=False, keep_links=True)
    report = {
        "path": str(path),
        "sheetnames": wb.sheetnames,
        "defined_names": [],
        "calculation": {
            "calcMode": getattr(wb.calculation, "calcMode", None),
            "fullCalcOnLoad": getattr(wb.calculation, "fullCalcOnLoad", None),
            "forceFullCalc": getattr(wb.calculation, "forceFullCalc", None),
        },
        "external_links": len(getattr(wb, "_external_links", [])),
        "sheets": [],
    }
    for name, obj in wb.defined_names.items():
        report["defined_names"].append({"name": name, "text": obj.attr_text, "hidden": obj.hidden})

    for ws in wb.worksheets:
        values_ws = wb_values[ws.title]
        formulas = []
        errors = []
        nonempty = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if len(nonempty) < 400:
                    nonempty.append({"cell": cell.coordinate, "value": cell.value, "number_format": cell.number_format})
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                        "cached": values_ws[cell.coordinate].value,
                    })
                if cell.data_type == "e" or (isinstance(values_ws[cell.coordinate].value, str) and values_ws[cell.coordinate].value.startswith("#")):
                    errors.append({"cell": cell.coordinate, "value": values_ws[cell.coordinate].value})
        report["sheets"].append({
            "title": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "hidden_rows": [idx for idx, dim in ws.row_dimensions.items() if dim.hidden],
            "hidden_columns": [idx for idx, dim in ws.column_dimensions.items() if dim.hidden],
            "sheet_protection": bool(ws.protection.sheet),
            "tables": [{"name": t.name, "ref": t.ref} for t in ws.tables.values()],
            "data_validations": len(ws.data_validations.dataValidation),
            "comments": [{"cell": c.coordinate, "text": c.comment.text} for row in ws.iter_rows() for c in row if c.comment],
            "formulas": formulas,
            "formula_errors": errors,
            "nonempty_sample": nonempty,
        })
    with ZipFile(path) as zf:
        report["zip_members"] = zf.namelist()
        report["vba"] = any(n.lower().endswith("vbaproject.bin") for n in zf.namelist())
        report["connections"] = any(n == "xl/connections.xml" for n in zf.namelist())
        report["external_link_parts"] = [n for n in zf.namelist() if n.startswith("xl/externalLinks/")]
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("files", nargs="+", type=Path)
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)
    for file in args.files:
        report = workbook_report(file)
        out = args.out_dir / f"{file.stem}-ooxml.json"
        out.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        print(out)


if __name__ == "__main__":
    main()
