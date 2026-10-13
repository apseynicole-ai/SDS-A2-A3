#!/usr/bin/env python3
"""Independent calculations for all eight altered-number Excel drills."""

from __future__ import annotations

import json
import math
from pathlib import Path
from statistics import NormalDist

import numpy as np
from openpyxl import load_workbook


STD_NORMAL = NormalDist()


def betacf(a: float, b: float, x: float) -> float:
    """Continued fraction for the incomplete beta function."""
    max_iter, eps, fpmin = 300, 3e-14, 1e-300
    qab, qap, qam = a + b, a + 1.0, a - 1.0
    c = 1.0
    d = 1.0 - qab * x / qap
    if abs(d) < fpmin:
        d = fpmin
    d = 1.0 / d
    h = d
    for m in range(1, max_iter + 1):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1.0 + aa * d
        if abs(d) < fpmin:
            d = fpmin
        c = 1.0 + aa / c
        if abs(c) < fpmin:
            c = fpmin
        d = 1.0 / d
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1.0 + aa * d
        if abs(d) < fpmin:
            d = fpmin
        c = 1.0 + aa / c
        if abs(c) < fpmin:
            c = fpmin
        d = 1.0 / d
        delta = d * c
        h *= delta
        if abs(delta - 1.0) < eps:
            return h
    raise RuntimeError("incomplete beta did not converge")


def regularized_beta(x: float, a: float, b: float) -> float:
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    bt = math.exp(math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log1p(-x))
    if x < (a + 1.0) / (a + b + 2.0):
        return bt * betacf(a, b, x) / a
    return 1.0 - bt * betacf(b, a, 1.0 - x) / b


def t_cdf(value: float, df: int) -> float:
    x = df / (df + value * value)
    tail = 0.5 * regularized_beta(x, df / 2.0, 0.5)
    return 1.0 - tail if value >= 0 else tail


def t_ppf(probability: float, df: int) -> float:
    lo, hi = -20.0, 20.0
    for _ in range(200):
        mid = (lo + hi) / 2
        if t_cdf(mid, df) < probability:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2


BOOK = Path("/Users/nicole/Downloads/SDS188_A1S2_2026_Excel_Practice_Companion.xlsx")
wb = load_workbook(BOOK, data_only=False, read_only=True)

# EX01
ws = wb["Bootstrap_Practice"]
resample_means = [
    float(np.mean([ws.cell(row=r, column=c).value for c in range(3, 13)]))
    for r in range(5, 205)
]
ordered = sorted(resample_means)

# EX02
p = 273 / 420
se = math.sqrt(p * (1 - p) / 420)
upper = p + STD_NORMAL.inv_cdf(0.975) * se

# EX03
tcrit = t_ppf(0.95, 19)

# EX04
calls = wb["Calls_Practice"]
call_values = [calls.cell(row=r, column=2).value for r in range(5, 41)]
n_calls = len(call_values)
xbar_calls = float(np.mean(call_values))
se_calls = 4.5 / math.sqrt(n_calls)
zcrit = STD_NORMAL.inv_cdf(0.99)
boundary = 12 + zcrit * se_calls
zcalc = (xbar_calls - 12) / se_calls
beta = STD_NORMAL.cdf((boundary - 14) / se_calls)
power = 1 - beta

# EX05
n0_no_prior = STD_NORMAL.inv_cdf(0.975) ** 2 * 0.5 * 0.5 / 0.055**2

# EX06
z99 = STD_NORMAL.inv_cdf(0.995)
n0_finite = z99**2 * 0.42 * 0.58 / 0.025**2
n_finite = n0_finite * 2400 / (n0_finite + 2399)

# EX08
tl = (21.5 - 22) / (5 / math.sqrt(30))
tu = (22.5 - 22) / (5 / math.sqrt(30))
tprob = t_cdf(tu, 29) - t_cdf(tl, 29)
zprob = STD_NORMAL.cdf(tu) - STD_NORMAL.cdf(tl)

result = {
    "EX01": {
        "count": len(resample_means),
        "minimum": ordered[0],
        "maximum": ordered[-1],
        "5th_smallest": ordered[4],
        "6th_smallest_after_cutting_5": ordered[5],
        "5th_largest": ordered[-5],
        "6th_largest_after_cutting_5": ordered[-6],
        "numpy_linear_2.5_percentile": float(np.percentile(resample_means, 2.5, method="linear")),
        "numpy_linear_97.5_percentile": float(np.percentile(resample_means, 97.5, method="linear")),
    },
    "EX02": {"p": p, "se": se, "upper": upper},
    "EX03": {"df": 19, "critical_t": tcrit, "decision_t_2.05": "Reject H0" if 2.05 > tcrit else "Do not reject H0"},
    "EX04": {
        "n": n_calls,
        "mean": xbar_calls,
        "se": se_calls,
        "critical_z": zcrit,
        "critical_boundary": boundary,
        "zcalc": zcalc,
        "beta_true_mean_14": beta,
        "power_true_mean_14": power,
    },
    "EX05": {"unrounded_n": n0_no_prior, "ceiling": math.ceil(n0_no_prior)},
    "EX06": {"unadjusted_n0": n0_finite, "adjusted_n": n_finite, "ceiling": math.ceil(n_finite)},
    "EX07": {"answer": "Do not reject H0 when H0 is false"},
    "EX08": {"t_lower": tl, "t_upper": tu, "df": 29, "t_probability": tprob, "normal_probability_trap": zprob},
}
print(json.dumps(result, indent=2))
