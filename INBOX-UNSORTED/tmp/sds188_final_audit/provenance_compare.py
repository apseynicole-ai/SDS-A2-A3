#!/usr/bin/env python3
"""Quantify lesson-text traceability from the final HTML into the offline DOCX."""

from __future__ import annotations

import html
import json
import re
import unicodedata
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
HTML_PATH = ROOT / "00_FINAL_USE/SDS188_A1S2_2026_MasterLearningSystem_FINAL.html"
DOCX_PATH = ROOT / "SDS188_A1S2_2026_COMPLETE_OFFLINE_PRINT_PACK (1).docx"


def norm(value: str) -> str:
    value = html.unescape(value)
    value = value.replace("\\n", " ").replace("\\u00a0", " ")
    value = unicodedata.normalize("NFKC", value)
    value = value.translate(str.maketrans({
        "‘": "'", "’": "'", "“": '"', "”": '"',
        "–": "-", "—": "-", "−": "-", "×": "x", "÷": "/",
    }))
    # Do not apply generic tag stripping here: statistical inequalities such
    # as X<c are valid text and can otherwise consume later prose.
    value = re.sub(r"[^0-9a-zA-Zα-ωΑ-Ω]+", " ", value.casefold())
    return " ".join(value.split())


html_text = norm(HTML_PATH.read_text(encoding="utf-8"))
doc = Document(DOCX_PATH)

paragraphs = [p.text.strip() for p in doc.paragraphs]
lesson_start = next(i for i, text in enumerate(paragraphs) if "LESSON 6.1" in text)
lesson_end = next(i for i, text in enumerate(paragraphs) if text == "CHAPTER 9 A1S2 STOP POINT")

blocks: list[dict[str, object]] = []
for index, value in enumerate(paragraphs[lesson_start:lesson_end], start=lesson_start):
    normalized = norm(value)
    if len(normalized) < 40:
        continue
    if "offline guided learn" in normalized or normalized.startswith("chapter "):
        continue
    matched = normalized in html_text
    blocks.append({
        "index": index,
        "chars": len(normalized),
        "matched": matched,
        "text": value[:220],
    })

matched = [b for b in blocks if b["matched"]]
unmatched = [b for b in blocks if not b["matched"]]
matched_chars = sum(int(b["chars"]) for b in matched)
total_chars = sum(int(b["chars"]) for b in blocks)

result = {
    "html": str(HTML_PATH),
    "docx": str(DOCX_PATH),
    "lesson_range_paragraphs": [lesson_start, lesson_end - 1],
    "substantive_lesson_blocks": len(blocks),
    "exact_normalized_substring_matches": len(matched),
    "unmatched_blocks": len(unmatched),
    "matched_character_share": matched_chars / total_chars if total_chars else None,
    "unmatched_examples": unmatched[:50],
}

# The HTML keeps the guided lessons as JSON-compatible JavaScript arrays.
# Compare the source fields to the whole DOCX corpus, allowing the DOCX's
# print labels ("Worked example", "Result", "Verify") to be additional text.
raw_html = HTML_PATH.read_text(encoding="utf-8")
docx_corpus = norm("\n".join(paragraphs + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]))
source_fields: list[dict[str, object]] = []


def visit(value: object, path: str) -> None:
    if isinstance(value, str):
        normalized = norm(value)
        if len(normalized) >= 20:
            source_fields.append({"path": path, "chars": len(normalized), "matched": normalized in docx_corpus, "text": value[:220]})
    elif isinstance(value, list):
        for i, item in enumerate(value):
            visit(item, f"{path}[{i}]")
    elif isinstance(value, dict):
        for key, item in value.items():
            if key in {"visual", "practiceId", "type", "tolerance"}:
                continue
            visit(item, f"{path}.{key}")


for chapter in (6, 7, 8, 9):
    match = re.search(rf"const CH{chapter}_LESSONS=(\[.*?\]);", raw_html, flags=re.S)
    if not match:
        raise RuntimeError(f"CH{chapter}_LESSONS not found")
    lessons = json.loads(match.group(1))
    if chapter == 9:
        lessons = [lesson for lesson in lessons if int(str(lesson["id"]).split(".", 1)[1]) <= 7]
    for lesson in lessons:
        visit(lesson, f"CH{chapter}.{lesson['id']}")

field_matches = [item for item in source_fields if item["matched"]]
field_missing = [item for item in source_fields if not item["matched"]]
result["html_lesson_field_comparison"] = {
    "included_lessons": {"6": 10, "7": 13, "8": 13, "9": 7},
    "substantive_source_fields": len(source_fields),
    "represented_in_docx": len(field_matches),
    "missing_from_docx": len(field_missing),
    "represented_character_share": (
        sum(int(item["chars"]) for item in field_matches) /
        sum(int(item["chars"]) for item in source_fields)
    ),
    "missing_examples": field_missing[:50],
}
print(json.dumps(result, indent=2, ensure_ascii=False))
